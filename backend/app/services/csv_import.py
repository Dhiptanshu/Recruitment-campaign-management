import csv
import io
import re

from sqlalchemy.orm import Session

from ..models import Candidate

PHONE_RE = re.compile(r"^\+?[0-9]{7,15}$")
EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
REQUIRED_COLUMNS = {"name", "phone"}
SUPPORTED_EXTENSIONS = {"csv", "xlsx", "xlsm"}
BATCH_SIZE = 500


class ImportSummary:
    def __init__(self):
        self.total_rows = 0
        self.imported = 0
        self.updated = 0
        self.skipped = 0
        self.errors: list[dict] = []  # {row, reason}
        self.touched_candidate_ids: list[int] = []

    def to_dict(self):
        return {
            "total_rows": self.total_rows,
            "imported": self.imported,
            "updated": self.updated,
            "skipped": self.skipped,
            "error_count": len(self.errors),
            "errors": self.errors[:50],
        }


def _clean(value):
    if value is None:
        return None
    v = str(value).strip()
    return v if v else None


def _normalize_phone(phone: str) -> str:
    return re.sub(r"[^0-9+]", "", phone)


# must match (or be tighter than) the column widths in models.py -- Postgres
# enforces VARCHAR(n) strictly (unlike SQLite, which silently accepts
# anything), so without this check a too-long field imports fine in local
# dev and then 500s in production.
MAX_NAME_LEN = 200
MAX_PHONE_LEN = 32
MAX_EMAIL_LEN = 200
MAX_COMPANY_LEN = 200
MAX_EXTERNAL_REF_LEN = 100


def validate_row(row: dict) -> tuple[dict | None, str | None]:
    name = _clean(row.get("name"))
    phone_raw = _clean(row.get("phone"))
    email = _clean(row.get("email"))
    company = _clean(row.get("current_company"))
    external_ref = _clean(row.get("id"))

    if not name:
        return None, "missing name"
    if len(name) > MAX_NAME_LEN:
        return None, f"name too long (max {MAX_NAME_LEN} characters)"
    if not phone_raw:
        return None, "missing phone"
    phone = _normalize_phone(phone_raw)
    if not PHONE_RE.match(phone):
        return None, f"invalid phone '{phone_raw}'"
    if len(phone) > MAX_PHONE_LEN:
        return None, f"phone too long (max {MAX_PHONE_LEN} characters)"
    if email:
        if not EMAIL_RE.match(email):
            return None, f"invalid email '{email}'"
        if len(email) > MAX_EMAIL_LEN:
            return None, f"email too long (max {MAX_EMAIL_LEN} characters)"
    if company and len(company) > MAX_COMPANY_LEN:
        company = company[:MAX_COMPANY_LEN]
    if external_ref and len(external_ref) > MAX_EXTERNAL_REF_LEN:
        external_ref = external_ref[:MAX_EXTERNAL_REF_LEN]

    return {
        "name": name,
        "phone": phone,
        "email": email,
        "current_company": company,
        "external_ref": external_ref,
    }, None


def _cell_to_str(value):
    """openpyxl hands back native ints/floats for numeric-looking cells
    (e.g. a phone column Excel decided was a number) -- normalize those
    back to plain digit strings instead of '919876543210.0'."""
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _iter_csv_rows(file_bytes: bytes):
    """Returns (header_set, row_iterator) or (None, None) if the file is empty.
    Each yielded item is (row_number, {lowercased_column: raw_value})."""
    try:
        text = file_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1")

    reader = csv.DictReader(io.StringIO(text))
    if reader.fieldnames is None:
        return None, None

    header = {h.strip().lower() for h in reader.fieldnames if h}

    def gen():
        for i, raw_row in enumerate(reader, start=2):  # header is row 1
            yield i, {(k or "").strip().lower(): v for k, v in raw_row.items()}

    return header, gen()


def _iter_xlsx_rows(file_bytes: bytes):
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_row = next(rows_iter)
    except StopIteration:
        return None, None

    headers = [(_cell_to_str(h) or "").lower() for h in header_row]
    header = {h for h in headers if h}

    def gen():
        for i, raw_row in enumerate(rows_iter, start=2):
            if all(v is None for v in raw_row):
                continue  # skip fully blank rows (common at the end of a sheet)
            row_dict = {
                headers[idx]: _cell_to_str(v)
                for idx, v in enumerate(raw_row)
                if idx < len(headers) and headers[idx]
            }
            yield i, row_dict

    return header, gen()


def import_candidates_file(db: Session, filename: str, file_bytes: bytes) -> ImportSummary:
    summary = ImportSummary()

    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "csv":
        header, rows = _iter_csv_rows(file_bytes)
    elif ext in ("xlsx", "xlsm"):
        header, rows = _iter_xlsx_rows(file_bytes)
    else:
        summary.errors.append({"row": 0, "reason": f"unsupported file type '.{ext}' -- use .csv or .xlsx"})
        return summary

    if header is None:
        summary.errors.append({"row": 0, "reason": "empty file"})
        return summary

    missing_cols = REQUIRED_COLUMNS - header
    if missing_cols:
        summary.errors.append({
            "row": 0,
            "reason": f"missing required column(s): {', '.join(sorted(missing_cols))}",
        })
        return summary

    # keyed by phone so duplicate phones within one flush cycle collapse to
    # the last occurrence instead of violating the unique constraint.
    batch: dict[str, dict] = {}

    def flush():
        nonlocal batch
        if not batch:
            return
        existing = {
            c.phone: c for c in db.query(Candidate).filter(Candidate.phone.in_(batch.keys())).all()
        }
        new_candidates = []
        for phone, rec in batch.items():
            existing_candidate = existing.get(phone)
            if existing_candidate:
                existing_candidate.name = rec["name"]
                existing_candidate.email = rec["email"] or existing_candidate.email
                existing_candidate.current_company = rec["current_company"] or existing_candidate.current_company
                summary.updated += 1
                summary.touched_candidate_ids.append(existing_candidate.id)
            else:
                candidate = Candidate(**rec)
                db.add(candidate)
                new_candidates.append(candidate)
                summary.imported += 1
        db.flush()
        summary.touched_candidate_ids.extend(c.id for c in new_candidates)
        db.commit()
        batch = {}

    for i, normalized_row in rows:
        summary.total_rows += 1
        record, error = validate_row(normalized_row)
        if error:
            summary.skipped += 1
            if len(summary.errors) < 200:
                summary.errors.append({"row": i, "reason": error})
            continue
        batch[record["phone"]] = record
        if len(batch) >= BATCH_SIZE:
            flush()

    flush()
    return summary
